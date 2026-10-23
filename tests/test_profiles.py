import datetime
import json
import os
import shutil
import tempfile
import unittest
import zipfile

from docx import Document
from openpyxl import Workbook

import app.repository.config_repository as repo
from app.core.profile_manager import (
    ProfileManager, ProfileError, slugify, MAPPING_FILE, TEMPLATE_FILE, DATA_FILE,
)
from app.core.report_generator import ReportGenerator
from app.core.state_manager import ReportStateManager
from app.gui.report_actions import generate_report


MAPPING = {
    "config": {"data_file": "S:/old/data.xlsx", "template_file": "S:/old/t.docx",
               "output_dir": "", "date_format": "%d/%m/%Y"},
    "sample_name": {"column": "sample name", "placeholder": "{{sample_name}}"},
    "coa_number":  {"column": "coa number",  "placeholder": "{{coa_number}}"},
    "file_name":   {"operation": "format", "format": "COA {coa_number}.docx"},
}


def make_files(folder, mapping=None, headers=("sample name", "coa number"), rows=(("Yeast", "C-1"),),
               sheet="Data", placeholders=("{{sample_name}}", "{{coa_number}}")):
    os.makedirs(folder, exist_ok=True)
    m = os.path.join(folder, "m.json")
    with open(m, "w", encoding="utf-8") as f:
        json.dump(mapping or MAPPING, f)
    d = os.path.join(folder, "d.xlsx")
    wb = Workbook(); ws = wb.active; ws.title = sheet
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    wb.save(d)
    t = os.path.join(folder, "t.docx")
    doc = Document()
    for p in placeholders:
        doc.add_paragraph(p)
    doc.save(t)
    return m, t, d


class ProfileTestBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self._old_db = repo.DB_PATH
        repo.DB_PATH = os.path.join(self.tmp, "app.db")
        self.pm = ProfileManager(root=os.path.join(self.tmp, "profiles"))
        self.src = os.path.join(self.tmp, "src")
        self.m, self.t, self.d = make_files(self.src)

    def tearDown(self):
        repo.DB_PATH = self._old_db
        shutil.rmtree(self.tmp, ignore_errors=True)


class TestSlug(unittest.TestCase):
    def test_slugify(self):
        self.assertEqual(slugify("Fish quality (v2)"), "fish-quality-v2")
        self.assertEqual(slugify("Café Été"), "cafe-ete")
        self.assertEqual(slugify("../..\\evil"), "evil")
        self.assertEqual(slugify("CON"), "con-profile")
        self.assertEqual(slugify("???"), "profile")


class TestProfileCrud(ProfileTestBase):
    def test_create_copies_files_and_strips_absolute_paths(self):
        p = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        self.assertEqual(p.slug, "coa")
        for f in (MAPPING_FILE, TEMPLATE_FILE, DATA_FILE):
            self.assertTrue((p.folder / f).is_file())
        cfg = json.load(open(p.mapping_path, encoding="utf-8"))["config"]
        self.assertNotIn("data_file", cfg)
        self.assertNotIn("template_file", cfg)
        self.assertEqual(cfg["profile_name"], "COA")
        # originals untouched
        self.assertIn("data_file", json.load(open(self.m))["config"])

    def test_create_without_files_makes_starter_mapping(self):
        p = self.pm.create("Empty")
        self.assertTrue(os.path.isfile(p.mapping_path))
        codes = [i.message for i in self.pm.validate(p.slug)]
        self.assertTrue(any("template.docx is missing" in c for c in codes))
        self.assertTrue(any("data.xlsx is missing" in c for c in codes))

    def test_slug_uniqueness_and_name_uniqueness(self):
        a = self.pm.create("COA")
        with self.assertRaises(ProfileError):
            self.pm.create("coa")          # same name, different case
        b = self.pm.create("C.O.A")        # different name, same slug base
        self.assertNotEqual(a.slug, b.slug)

    def test_bad_names(self):
        for bad in ("", "   "):
            with self.assertRaises(ProfileError):
                self.pm.create(bad)
        p = self.pm.create("../../escape")
        self.assertEqual(p.folder.parent, self.pm.root)   # stayed inside the root

    def test_failed_create_leaves_nothing(self):
        with self.assertRaises(ProfileError):
            self.pm.create("Broken", mapping=os.path.join(self.tmp, "nope.json"))
        self.assertEqual(self.pm.list(), [])
        self.assertEqual(os.listdir(self.pm.root), [])

    def test_duplicate_rename_delete(self):
        a = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        b = self.pm.duplicate(a.slug, "COA copy")
        self.assertEqual(json.load(open(b.mapping_path))["sample_name"]["column"], "sample name")
        self.assertEqual(json.load(open(b.mapping_path))["config"]["profile_name"], "COA copy")
        self.assertTrue((b.folder / DATA_FILE).is_file())

        c = self.pm.rename(b.slug, "Renamed")
        self.assertEqual(c.name, "Renamed")
        self.assertEqual(c.slug, b.slug)       # folder unchanged
        with self.assertRaises(ProfileError):
            self.pm.rename(b.slug, "COA")      # clashes with a

        trashed = self.pm.delete(b.slug)
        self.assertFalse(b.folder.exists())
        self.assertTrue(trashed.is_dir())      # moved, not destroyed
        self.assertEqual([p.slug for p in self.pm.list()], ["coa"])

    def test_create_from_existing_with_overrides(self):
        a = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        _, _, d2 = make_files(os.path.join(self.tmp, "src9"), rows=(("Other", "C-9"),))
        b = self.pm.create_from("COA v2", source_slug=a.slug, data=d2)
        self.assertEqual(json.load(open(b.mapping_path))["config"]["profile_name"], "COA v2")
        self.assertTrue((b.folder / "data.xlsx.bak").is_file())      # replaced the copied data
        with self.assertRaises(ProfileError):
            self.pm.create_from("Bad", source_slug=a.slug, template=self.d)   # wrong extension
        self.assertIsNone(self.pm.find_by_name("Bad"))                        # rolled back
        self.assertEqual(len(self.pm.list()), 2)

    def test_deleting_active_profile_clears_active(self):
        a = self.pm.create("A"); b = self.pm.create("B")
        self.pm.set_active(a.slug)
        self.assertEqual(self.pm.active_slug(), "a")
        self.pm.delete(a.slug)
        self.assertIsNone(self.pm.active_slug())

    def test_scan_registers_hand_copied_folder_and_prunes_missing(self):
        a = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        shutil.copytree(a.folder, self.pm.root / "manual")
        names = sorted(p.name for p in self.pm.list())
        self.assertEqual(len(names), 2)                    # "COA" and "COA (2)" (same profile_name)
        shutil.rmtree(a.folder)
        self.assertEqual([p.slug for p in self.pm.list()], ["manual"])

    def test_replace_data_makes_backup(self):
        p = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        _, _, d2 = make_files(os.path.join(self.tmp, "src2"), rows=(("Other", "C-9"),))
        self.pm.replace_data(p.slug, d2)
        self.assertTrue((p.folder / "data.xlsx.bak").is_file())
        with self.assertRaises(ProfileError):
            self.pm.replace_data(p.slug, self.t)           # wrong extension

    def test_next_line_and_output_dir(self):
        p = self.pm.create("COA")
        self.assertEqual(p.next_line, 2)
        self.pm.set_next_line(p.slug, 7)
        self.pm.set_next_line(p.slug, 1)                   # clamped to header+1
        self.assertEqual(self.pm.get(p.slug).next_line, 2)
        self.pm.set_next_line(p.slug, 7)
        self.pm.set_output_dir(p.slug, "/x/y")
        got = self.pm.get(p.slug)
        self.assertEqual((got.next_line, got.output_dir), (7, "/x/y"))


class TestValidate(ProfileTestBase):
    def test_ok_profile_has_no_issues(self):
        p = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        self.assertEqual(self.pm.validate(p.slug), [])

    def test_invalid_json(self):
        p = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        open(p.mapping_path, "w").write("{ not json")
        issues = self.pm.validate(p.slug)
        self.assertTrue(any(i.severity == "error" and "not valid" in i.message for i in issues))

    def test_missing_data_sheet(self):
        mp = dict(MAPPING); mp["config"] = {**MAPPING["config"], "data_sheet": "Nope"}
        m, t, d = make_files(os.path.join(self.tmp, "s3"), mapping=mp)
        p = self.pm.create("X", mapping=m, template=t, data=d)
        issues = self.pm.validate(p.slug)
        self.assertTrue(any(i.severity == "error" and "Nope" in i.message for i in issues))

    def test_data_sheet_case_insensitive(self):
        mp = dict(MAPPING); mp["config"] = {**MAPPING["config"], "data_sheet": "data"}
        m, t, d = make_files(os.path.join(self.tmp, "s4"), mapping=mp)
        p = self.pm.create("X", mapping=m, template=t, data=d)
        self.assertEqual(self.pm.validate(p.slug), [])

    def test_column_and_placeholder_warnings(self):
        m, t, d = make_files(os.path.join(self.tmp, "s5"),
                             headers=("sample name",), placeholders=("{{sample_name}}",))
        p = self.pm.create("X", mapping=m, template=t, data=d)
        msgs = [(i.severity, i.message) for i in self.pm.validate(p.slug)]
        self.assertTrue(any(s == "warning" and "column 'coa number'" in m for s, m in msgs))
        self.assertTrue(any(s == "warning" and "{{coa_number}}" in m for s, m in msgs))
        self.assertFalse(any(s == "error" for s, _ in msgs))


def make_lookup_files(folder, mapping, data_rows, suppliers, headers=("fournisseur fabricant", "espece", "K-value", "agrement"),
                      placeholders=("{{agrement}}", "{{k_value}}")):
    os.makedirs(folder, exist_ok=True)
    m = os.path.join(folder, "m.json"); json.dump(mapping, open(m, "w", encoding="utf-8"))
    wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(list(headers))
    for r in data_rows:
        ws.append(list(r))
    f = wb.create_sheet("fournisseurs"); f.append(["Fournisseur", "Numero agrement"])
    for r in suppliers:
        f.append(list(r))
    e = wb.create_sheet("especes"); e.append(["Nom Latin", "Nom Français"])
    e.append(["Thunnus albacares", "Thon albacore"])
    c = wb.create_sheet("crit"); c.append(["Espece", "Critere"]); c.append(["Thon albacore", "Crit A"])
    d = os.path.join(folder, "d.xlsx"); wb.save(d)
    t = os.path.join(folder, "t.docx"); doc = Document()
    for ph in placeholders:
        doc.add_paragraph(ph)
    doc.save(t)
    return m, t, d


LOOKUP = {"config": {"data_sheet": "Data"},
          "agrement": {"operation": "lookup", "sheet": "fournisseurs", "key": "Fournisseur",
                       "match": "fournisseur fabricant", "value": "Numero agrement",
                       "placeholder": "{{agrement}}"},
          "k-value": {"column": "K-value", "placeholder": "{{k_value}}"}}


class TestChecks(ProfileTestBase):
    def _profile(self, mapping, rows, suppliers=(("ACME", "FR1"),), **kw):
        m, t, d = make_lookup_files(os.path.join(self.tmp, "chk"), mapping, rows, suppliers, **kw)
        return self.pm.create("Chk", mapping=m, template=t, data=d)

    def msgs(self, p, **kw):
        return [i.message for i in self.pm.validate(p.slug, **kw)]

    def test_clean_lookup_has_no_issues_and_matches_like_the_generator(self):
        # generator compares stripped + case-insensitive, so these all resolve
        p = self._profile(LOOKUP, [("acme", "x", 1, ""), ("ACME\xa0", "x", 2, ""), (" Acme ", "x", 3, "")])
        self.assertEqual(self.msgs(p), [])

    def test_missing_lookup_values_reported_with_rows_and_fallback_hint(self):
        p = self._profile(LOOKUP, [("ACME", "x", 1, ""), ("big fish", "x", 2, "FAO 51"),
                                   ("big fish", "x", 3, ""), ("other", "x", 4, "")])
        found = [m for m in self.msgs(p) if "Field 'agrement'" in m]
        self.assertEqual(len(found), 1)
        msg = found[0]
        self.assertIn("2 value(s) of 'fournisseur fabricant'", msg)
        self.assertIn("{{agrement}} will be empty for 3 row(s)", msg)
        self.assertIn("'big fish' (rows 3, 4)", msg)
        self.assertIn("own 'agrement' column", msg)
        self.assertEqual(self.msgs(p, deep=False), [])             # quick check skips data rows

    def test_did_you_mean_for_near_miss_supplier(self):
        p = self._profile(LOOKUP, [("Acme Seafood", "x", 1, "")], suppliers=(("ACME Seafod", "FR1"),))
        self.assertTrue(any("did you mean 'ACME Seafod'" in m for m in self.msgs(p)))

    def test_found_but_blank_value(self):
        p = self._profile(LOOKUP, [("ACME", "x", 1, "")], suppliers=(("ACME", None),))
        self.assertTrue(any("found but the looked-up cell is empty" in m for m in self.msgs(p)))

    def test_empty_input_cells_are_not_reported(self):
        p = self._profile(LOOKUP, [(None, "x", 1, ""), ("", "x", 2, "")])
        self.assertEqual(self.msgs(p), [])

    def test_column_hint_k_value(self):
        mp = {**LOOKUP, "k-value": {"column": "k value", "placeholder": "{{k_value}}"}}
        p = self._profile(mp, [("ACME", "x", 1, "")])
        msg = [m for m in self.msgs(p) if "k value" in m][0]
        self.assertIn("Did you mean 'K-value'?", msg)

    def test_template_placeholder_without_mapping(self):
        p = self._profile(LOOKUP, [("ACME", "x", 1, "")], placeholders=("{{agrement}}", "{{k_value}}", "{{n_ech}}"))
        self.assertEqual([m for m in self.msgs(p) if "n_ech" in m],
                         ["Template placeholder {{n_ech}} is not filled by any field in the mapping "
                          "— it will stay as {{n_ech}} in the report."])

    def test_computed_field_without_placeholder_key_is_covered_by_its_key(self):
        mp = {**LOOKUP, "date_du_jour": {"operation": "today", "format": "%d/%m/%Y"},
              "numero_rapport": {"operation": "report_number"}}
        p = self._profile(mp, [("ACME", "x", 1, "")],
                          placeholders=("{{agrement}}", "{{k_value}}", "{{date_du_jour}}", "{{numéro_rapport}}"))
        self.assertEqual(self.msgs(p), [])

    def test_lookup_structure_errors(self):
        bad = {**LOOKUP, "agrement": {**LOOKUP["agrement"], "sheet": "fournisseur", "value": "x"}}
        p = self._profile(bad, [("ACME", "x", 1, "")])
        m = self.msgs(p)
        self.assertTrue(any("sheet 'fournisseur' not found" in x and "Did you mean 'fournisseurs'" in x for x in m))
        bad2 = {**LOOKUP, "agrement": {**LOOKUP["agrement"], "key": "fournisseur", "match": "fournisseur"}}
        p2 = self.pm.replace_mapping(p.slug, self._write(bad2)) or p
        m2 = self.msgs(p2)
        self.assertTrue(any("column 'fournisseur' not found in sheet 'fournisseurs'" in x for x in m2))
        self.assertTrue(any("match column 'fournisseur' not found in the data sheet" in x for x in m2))

    def _write(self, mapping):
        path = os.path.join(self.tmp, "tmp_map.json"); json.dump(mapping, open(path, "w")); return path

    def test_lookup_join_chain(self):
        join = {"config": {"data_sheet": "Data"},
                "crit": {"operation": "lookup_join", "placeholder": "{{agrement}}", "steps": [
                    {"sheet": "especes", "key": "Nom Latin", "match": "espece", "value": "Nom Français"},
                    {"sheet": "crit", "key": "Espece", "match": "<previous>", "value": "Critere"}]}}
        ok = self._profile(join, [("a", "Thunnus albacares", 1, "")], placeholders=("{{agrement}}",))
        self.assertEqual(self.msgs(ok), [])
        self.pm.delete(ok.slug)
        p = self._profile(join, [("a", "Thunnus albacares", 1, ""), ("a", "Thunnus obesus", 2, "")],
                          placeholders=("{{agrement}}",))
        m = [x for x in self.msgs(p) if "Field 'crit'" in x]
        self.assertEqual(len(m), 1)
        self.assertIn("'Thunnus obesus' (row 3)", m[0])


class TestExportImport(ProfileTestBase):
    def test_roundtrip(self):
        p = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        z = os.path.join(self.tmp, "coa.zip")
        self.pm.export_zip(p.slug, z)
        q = self.pm.import_zip(z)
        self.assertNotEqual(q.slug, p.slug)
        self.assertEqual(q.name, "COA (2)")                # name clash resolved
        self.assertEqual(self.pm.validate(q.slug), [])

    def test_import_ignores_paths_inside_archive(self):
        z = os.path.join(self.tmp, "evil.zip")
        with zipfile.ZipFile(z, "w") as zf:
            zf.write(self.m, "../../mapping.json")
            zf.write(self.m, "other.txt")
        q = self.pm.import_zip(z, name="Evil")
        self.assertEqual(q.folder.parent, self.pm.root)
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "mapping.json")))

    def test_import_requires_mapping(self):
        z = os.path.join(self.tmp, "bad.zip")
        with zipfile.ZipFile(z, "w") as zf:
            zf.write(self.t, "template.docx")
        with self.assertRaises(ProfileError):
            self.pm.import_zip(z)


class TestMigration(ProfileTestBase):
    def _legacy(self, xlsx_mapping=False):
        # point the mapping's config at real files, like the old absolute-path setup
        cfg = {**MAPPING["config"], "data_file": self.d, "template_file": self.t,
               "output_dir": os.path.join(self.tmp, "out")}
        mp = {**MAPPING, "config": cfg}
        m = os.path.join(self.src, "coa_mapping.json")
        json.dump(mp, open(m, "w"))
        repo.mapping_add(m)
        repo.mapping_add(os.path.join(self.tmp, "gone.json"))      # registered but deleted
        repo.config_set("mapping_path", m)
        repo.config_set("last_line_number", "11")
        repo.config_set("last_output_dir", os.path.join(self.tmp, "last_out"))
        return m

    def test_migrates_once_and_keeps_originals(self):
        m = self._legacy()
        before = open(m).read()
        res = self.pm.migrate_legacy()
        self.assertEqual(res["created"], ["coa_mapping"])
        self.assertEqual(len(res["skipped"]), 1)
        p = self.pm.get(self.pm.active_slug())
        self.assertEqual(p.next_line, 12)
        self.assertTrue(p.output_dir.endswith("last_out"))
        self.assertEqual(self.pm.validate(p.slug), [])
        self.assertEqual(open(m).read(), before)                    # original untouched
        # second run does nothing
        self.assertEqual(self.pm.migrate_legacy(), {"created": [], "skipped": []})
        self.assertEqual(len(self.pm.list()), 1)

    def test_missing_files_still_create_flagged_profile(self):
        m = self._legacy()
        os.remove(self.d)
        res = self.pm.migrate_legacy()
        self.assertEqual(len(res["created"]), 1)
        slug = self.pm.active_slug()
        self.assertTrue(any("data.xlsx is missing" in i.message for i in self.pm.validate(slug)))

    def test_active_profile_falls_back_to_global_last_paths(self):
        m = self._legacy()
        cfg = json.load(open(m)); cfg["config"].pop("data_file"); cfg["config"].pop("template_file")
        json.dump(cfg, open(m, "w"))
        repo.config_set("last_excel_path", self.d)
        repo.config_set("last_docx_path", self.t)
        self.pm.migrate_legacy()
        self.assertEqual(self.pm.validate(self.pm.active_slug()), [])

    def test_xlsx_mapping_is_converted_to_json(self):
        wb = Workbook(); ws = wb.active; ws.title = "mappings"
        ws.append(["Spreadsheet Column", "Placeholder", "Operation", "Type", "Notes"])
        ws.append(["sample name", "{{sample_name}}", "", "text (input)", ""])
        x = os.path.join(self.src, "map.xlsx"); wb.save(x)
        repo.mapping_add(x)
        self.pm.migrate_legacy()
        p = self.pm.list()[0]
        self.assertTrue(json.load(open(p.mapping_path)))            # valid JSON


class TestPerProfileState(ProfileTestBase):
    def test_counters_are_independent(self):
        a = ReportStateManager(os.path.join(self.tmp, "x.json"), scope="a")
        b = ReportStateManager(os.path.join(self.tmp, "x.json"), scope="b")
        g = ReportStateManager(os.path.join(self.tmp, "x.json"))
        prefix = datetime.date.today().strftime("%y%m%d")
        self.assertEqual(a.generate_report_number(), f"{prefix}-01")
        self.assertEqual(a.generate_report_number(), f"{prefix}-02")
        self.assertEqual(b.generate_report_number(), f"{prefix}-01")
        self.assertEqual(g.get_current_counter(), 0)
        self.assertEqual(a.get_current_counter(), 2)
        a.set_counter(9)
        self.assertEqual(a.get_current_counter(), 9)
        self.assertEqual(b.get_current_counter(), 1)

    def test_generate_through_profile(self):
        p = self.pm.create("COA", mapping=self.m, template=self.t, data=self.d)
        out = os.path.join(self.tmp, "out.docx")
        generate_report(p.data_path, p.template_path, p.mapping_path, 2, out, profile_slug=p.slug)
        text = "\n".join(par.text for par in Document(out).paragraphs)
        self.assertIn("Yeast", text); self.assertIn("C-1", text)


if __name__ == "__main__":
    unittest.main()
