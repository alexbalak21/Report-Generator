"""
Keeps the in-app guide honest: its example must really produce what it claims,
its operation lists must match the engine, and the AI export must not leak data.
"""
import datetime
import json
import os
import re
import shutil
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook

import app.repository.config_repository as repo
from app.core import guide
from app.core.profile_checks import COLUMN_OPERATIONS, COMPUTED_OPERATIONS
from app.core.profile_manager import ProfileManager
from app.core.report_generator import ReportGenerator
from app.gui.report_actions import generate_report, resolve_output_path


def _guide_text():
    return guide.load_guide()


def _example_mapping(text):
    """The JSON block right after the <!-- test:example-mapping --> marker (raw guide file)."""
    from app.utils.paths import get_app_resource_path
    raw = open(get_app_resource_path(guide.GUIDE_FILE), encoding="utf-8").read()
    after = raw.split("<!-- test:example-mapping -->", 1)[1]
    return json.loads(re.search(r"```json\n(.*?)```", after, re.S).group(1))


def _backticked_names(section_text, first_column_only=True):
    names = []
    for line in section_text.splitlines():
        if line.startswith("|") and not line.startswith("|---"):
            cell = line.split("|")[1]
            names += re.findall(r"`([a-z_]+)`", cell)
    return names


class GuideBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self._old_db = repo.DB_PATH
        repo.DB_PATH = os.path.join(self.tmp, "app.db")
        self.pm = ProfileManager(root=os.path.join(self.tmp, "profiles"))

    def tearDown(self):
        repo.DB_PATH = self._old_db
        shutil.rmtree(self.tmp, ignore_errors=True)


class TestGuideFile(unittest.TestCase):
    def test_loads_without_maintenance_comments(self):
        text = _guide_text()
        self.assertNotIn("<!--", text)
        self.assertTrue(text.startswith("# Report Generator"))

    def test_all_json_blocks_are_valid_json(self):
        blocks = re.findall(r"```json\n(.*?)```", _guide_text(), re.S)
        self.assertGreaterEqual(len(blocks), 6)
        for block in blocks:
            snippet = block.strip()
            # fragments ("operations": [...] / "field": {...}) are wrapped to be parsed
            candidate = snippet if snippet.startswith("{") else "{" + snippet + "}"
            json.loads(candidate)

    def test_headings_for_the_jump_list(self):
        titles = [t for t, _ in guide.headings(_guide_text(), level=2)]
        self.assertEqual(titles[0].split(".")[0], "0")
        self.assertEqual(len(titles), len(set(titles)))
        self.assertIn("9. Messages of the automatic check (status line → Details…)", titles)

    def test_headings_ignore_code_blocks(self):
        text = "## A\n```\n## not a heading\n```\n## B\n"
        self.assertEqual([t for t, _ in guide.headings(text)], ["A", "B"])

    def test_operation_lists_match_the_engine_and_the_checks(self):
        text = _guide_text()
        sec74 = text.split("### 7.4", 1)[1].split("### 7.5", 1)[0]
        sec75 = text.split("### 7.5", 1)[1].split("### 7.6", 1)[0]
        self.assertEqual(set(_backticked_names(sec74)) - {"type"}, set(COLUMN_OPERATIONS))
        self.assertEqual(set(_backticked_names(sec75)) - {"operation"}, set(COMPUTED_OPERATIONS))
        # ...and the calculated operations are exactly what the generator implements
        impl = ReportGenerator.__new__(ReportGenerator)
        impl.state_manager = None
        self.assertEqual(set(ReportGenerator._build_operations(impl, None, "", 0)), set(COMPUTED_OPERATIONS))


class TestWorkedExample(GuideBase):
    """Section 11 of the guide, executed."""

    def _build(self):
        folder = os.path.join(self.tmp, "ex"); os.makedirs(folder)
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws.append(["sample name", "reception date", "weight", "supplier", "address"])
        ws.append(["Sample A1", datetime.date(2026, 3, 4), 12.3456, "acme seafood",
                   "12 Quai Gambetta, 62200 Boulogne-sur-Mer"])
        s = wb.create_sheet("suppliers"); s.append(["Supplier", "Approval number"])
        s.append(["ACME Seafood", "FR 62.001.001 CE"])
        d = os.path.join(folder, "d.xlsx"); wb.save(d)
        doc = Document()
        for line in ["CERTIFICATE {{title}}", "Sample: {{sample_name}}", "Received on: {{reception_date}}",
                     "Weight: {{weight}} g", "Supplier approval: {{approval}}", "Address:", "{{address}}",
                     "Issued on {{report_date}}"]:
            doc.add_paragraph(line)
        t = os.path.join(folder, "t.docx"); doc.save(t)
        m = os.path.join(folder, "m.json"); json.dump(_example_mapping(_guide_text()), open(m, "w"))
        return self.pm.create("Example", mapping=m, template=t, data=d)

    def test_example_is_clean_and_produces_the_documented_report(self):
        p = self._build()
        self.assertEqual([i.message for i in self.pm.validate(p.slug)], [])      # no warning on the guide's own example

        folder, name = resolve_output_path(p.mapping_path, p.data_path, p.template_path, 2, profile_slug=p.slug)
        self.assertEqual(name, "Certificate Sample A1.docx")

        out = os.path.join(self.tmp, "out.docx")
        generate_report(p.data_path, p.template_path, p.mapping_path, 2, out, profile_slug=p.slug)
        lines = []
        for para in Document(out).paragraphs:
            lines += para.text.split("\n")        # a Word line break comes back as \n
        today = datetime.date.today().strftime("%d/%m/%Y")
        self.assertEqual(lines, [
            "CERTIFICATE COA - Sample A1",
            "Sample: Sample A1",
            "Received on: 04/03/2026",
            "Weight: 12,35 g",
            "Supplier approval: FR 62.001.001 CE",
            "Address:",
            "12 Quai Gambetta",
            "62200 Boulogne-sur-Mer",
            f"Issued on {today}",
        ])


class TestAiExport(GuideBase):
    def _profile(self):
        folder = os.path.join(self.tmp, "s"); os.makedirs(folder)
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws.append(["sample name", "secret column"])
        ws.append(["TOP-SECRET-VALUE", "CONFIDENTIAL-123"])
        ws.append(["another", "x"])
        d = os.path.join(folder, "d.xlsx"); wb.save(d)
        doc = Document(); doc.add_paragraph("{{sample_name}} {{orphan}}")
        t = os.path.join(folder, "t.docx"); doc.save(t)
        m = os.path.join(folder, "m.json")
        json.dump({"config": {"data_file": "D:/Private/Secret/data.xlsx", "template_file": "D:/x.docx",
                              "output_dir": "D:/Clients/Acme", "date_format": "%d/%m/%Y"},
                   "sample_name": {"column": "sample name", "placeholder": "{{sample_name}}"}},
                  open(m, "w"))
        return self.pm.create("Ctx", mapping=m, template=t, data=d)

    def test_context_has_what_an_ai_needs_and_nothing_private(self):
        p = self._profile()
        issues = self.pm.validate(p.slug)
        ctx = guide.build_profile_context(p, issues)
        for expected in ("Profile name: Ctx", "```json", '"sample name"', "{{orphan}}", "{{sample_name}}",
                         'Sheet "Data" (2 data row(s))', '"secret column"',
                         "Template placeholder {{orphan}} is not filled"):
            self.assertIn(expected, ctx)
        for leaked in ("TOP-SECRET-VALUE", "CONFIDENTIAL-123", "D:/Private", "D:/Clients", "D:/x.docx"):
            self.assertNotIn(leaked, ctx)
        self.assertIn("(path hidden)", ctx)

    def test_context_survives_broken_files(self):
        p = self._profile()
        open(p.mapping_path, "w").write("{ broken")
        os.remove(p.template_path)
        ctx = guide.build_profile_context(p, self.pm.validate(p.slug))
        self.assertIn("not valid JSON", ctx)
        self.assertIn("{ broken", ctx)
        self.assertIn("Could not read template.docx", ctx)

    def test_ai_message_layout(self):
        msg = guide.build_ai_message("GUIDE BODY", "CTX BODY")
        self.assertLess(msg.index("Report Generator"), msg.index("GUIDE BODY"))
        self.assertLess(msg.index("GUIDE BODY"), msg.index("CTX BODY"))
        self.assertTrue(msg.rstrip().endswith("My question / problem:"))
        self.assertNotIn("CTX", guide.build_ai_message("GUIDE BODY"))


if __name__ == "__main__":
    unittest.main()
