"""
Checks that catch mapping mistakes before a report is generated.

Everything here returns lists of Issue(severity, message). No tkinter.
"""
from __future__ import annotations

import difflib
import re
from dataclasses import dataclass

_PLACEHOLDER_RE = re.compile(r"\{\{(.*?)\}\}")
LOOKUP_OPS = ("lookup", "lookup_join")

# Placeholders the generator fills itself when the mapping defines a report number.
_REPORT_NUMBER_KEYS = ("sample_number", "numero_rapport")
_REPORT_NUMBER_PLACEHOLDERS = {"numéro_rapport", "sample_number"}


@dataclass
class Issue:
    severity: str   # "error" | "warning"
    message: str


# ── helpers ────────────────────────────────────────────────────────────────────

def placeholder_name(value) -> str | None:
    """'{{end_date}}' -> 'end_date' (None if it is not a placeholder)."""
    if not isinstance(value, str):
        return None
    match = _PLACEHOLDER_RE.search(value)
    return match.group(1).strip() if match else None


def _squash(text) -> str:
    return re.sub(r"[\s_\-]+", "", str(text)).lower()


def suggest_name(name: str, candidates) -> str | None:
    """Closest existing name: same letters ignoring case/space/-/_ first, then fuzzy."""
    candidates = [str(c) for c in candidates]
    target = _squash(name)
    for cand in sorted(candidates):
        if _squash(cand) == target and cand != name:
            return cand
    by_lower = {c.lower(): c for c in candidates}
    close = difflib.get_close_matches(str(name).lower(), list(by_lower), n=1, cutoff=0.75)
    return by_lower[close[0]] if close and by_lower[close[0]] != name else None


def _did_you_mean(name: str, candidates) -> str:
    suggestion = suggest_name(name, candidates)
    if not suggestion:
        return ""
    return (f" Did you mean '{suggestion}'? "
            "(names must match the Excel header exactly, including capitals, spaces and hyphens)")


def read_all_headers(data_path: str) -> dict[str, set[str]]:
    """{sheet name: set of header names} without loading the whole workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(data_path, read_only=True, data_only=True)
    try:
        result = {}
        for name in wb.sheetnames:
            first = next(wb[name].iter_rows(min_row=1, max_row=1, values_only=True), ())
            result[name] = {str(c).strip() for c in first if c is not None}
        return result
    finally:
        wb.close()


def _lookup_rules(mapping: dict):
    return [(k, r) for k, r in mapping.items()
            if isinstance(r, dict) and r.get("operation") in LOOKUP_OPS]


def _steps(rule: dict) -> list[dict]:
    if rule.get("operation") == "lookup_join":
        return [s for s in rule.get("steps", []) if isinstance(s, dict)]
    return [rule]


def _norm(value) -> str:
    return str(value).strip() if value is not None else ""


# ── 1. template placeholders nothing fills ─────────────────────────────────────

def covered_placeholders(mapping: dict) -> set[str]:
    """Every placeholder name the generator can fill from this mapping."""
    covered: set[str] = set()
    for key, rule in mapping.items():
        if key == "file_name":
            continue
        if isinstance(rule, str):                       # legacy: "column": "{{placeholder}}"
            name = placeholder_name(rule)
            if name:
                covered.add(name)
        elif isinstance(rule, dict):
            name = placeholder_name(rule.get("placeholder"))
            if name:
                covered.add(name)
            elif "operation" in rule:                   # computed field: {{<key>}}
                covered.add(str(key))
    if any(k in mapping for k in _REPORT_NUMBER_KEYS):
        covered |= _REPORT_NUMBER_PLACEHOLDERS
    return covered


def check_placeholder_coverage(mapping: dict, template_placeholders: set[str]) -> list[Issue]:
    return [
        Issue("warning",
              f"Template placeholder {{{{{name}}}}} is not filled by any field in the mapping "
              f"— it will stay as {{{{{name}}}}} in the report.")
        for name in sorted(template_placeholders - covered_placeholders(mapping))
    ]


# ── 2. lookup rules: sheets / columns exist ────────────────────────────────────

def check_lookup_structure(mapping: dict, data_headers: set[str],
                           sheet_headers: dict[str, set[str]]) -> list[Issue]:
    issues: list[Issue] = []
    for key, rule in _lookup_rules(mapping):
        steps = _steps(rule)
        if not steps:
            issues.append(Issue("warning", f"Field '{key}': lookup has no steps."))
            continue
        for number, step in enumerate(steps, start=1):
            where = f"Field '{key}'" + (f" (step {number})" if len(steps) > 1 else "")
            sheet, key_col, match = step.get("sheet"), step.get("key"), step.get("match")
            value = step.get("value")
            if not all([sheet, key_col, match, value]):
                issues.append(Issue("warning", f"{where}: lookup is incomplete "
                                               "(needs sheet, key, match and value)."))
                continue
            if sheet not in sheet_headers:
                issues.append(Issue(
                    "warning", f"{where}: lookup sheet '{sheet}' not found in the data file."
                               + _did_you_mean(sheet, sheet_headers)))
                continue
            columns = sheet_headers[sheet]
            if key_col not in columns:
                issues.append(Issue(
                    "warning", f"{where}: column '{key_col}' not found in sheet '{sheet}'."
                               + _did_you_mean(key_col, columns)))
            for v in (value if isinstance(value, list) else [value]):
                if v not in columns:
                    issues.append(Issue(
                        "warning", f"{where}: column '{v}' not found in sheet '{sheet}'."
                                   + _did_you_mean(v, columns)))
            if match != "<previous>" and match not in data_headers:
                issues.append(Issue(
                    "warning", f"{where}: match column '{match}' not found in the data sheet."
                               + _did_you_mean(match, data_headers)))
    return issues


# ── 3. lookup rules: values in the data that find nothing ──────────────────────

def check_lookup_data(mapping: dict, config: dict, data_path: str,
                      max_examples: int = 5) -> list[Issue]:
    """
    Run each lookup rule (with the real generator code) over every data row and
    report the input values that give no result — the case where a placeholder
    silently comes out empty, e.g. a supplier missing from the 'fournisseurs' sheet.
    """
    rules = _lookup_rules(mapping)
    if not rules:
        return []

    from app.core import processors
    from app.core.excel_reader import ExcelReader

    excel = ExcelReader(data_path)
    excel.set_data_sheet(config.get("data_sheet"))
    excel.load()
    rows = excel.get_sheet(excel._active_sheet_name())
    data_headers = set(rows[0].keys()) if rows else set()

    issues: list[Issue] = []
    for key, rule in rules:
        steps = _steps(rule)
        fields = [s.get("match") for s in steps
                  if s.get("match") and s.get("match") != "<previous>"]
        if not fields:
            continue
        primary = fields[0]
        func = processors.op_lookup if rule.get("operation") == "lookup" else processors.op_lookup_join

        inputs: dict[tuple, list[int]] = {}
        for index, row in enumerate(rows):
            values = tuple(_norm(row.get(f)) for f in fields)
            if values[0]:                                # an empty cell is the user's choice
                inputs.setdefault(values, []).append(index + 2)

        not_found: list[tuple[str, list[int]]] = []
        blank: list[tuple[str, list[int]]] = []
        try:
            for values, row_numbers in inputs.items():
                result = func(rule, dict(zip(fields, values)), excel)
                if result is None:
                    not_found.append((values[0], row_numbers))
                elif result == "":
                    blank.append((values[0], row_numbers))
        except LookupError:
            continue                                    # sheet/column problems: reported by the structure check

        placeholder = placeholder_name(rule.get("placeholder")) or key
        first = steps[0]
        fallback = (f" The data sheet also has its own '{key}' column you could use instead."
                    if key in data_headers else "")

        if not_found:
            sheet_keys = []
            try:
                sheet_keys = [_norm(r.get(first.get("key"))) for r in excel.get_sheet(first.get("sheet"))]
            except KeyError:
                pass
            issues.append(Issue(
                "warning",
                f"Field '{key}': {len(not_found)} value(s) of '{primary}' have no match in sheet "
                f"'{first.get('sheet')}' — {{{{{placeholder}}}}} will be empty for "
                f"{sum(len(r) for _, r in not_found)} row(s). "
                f"e.g. {_examples(not_found, sheet_keys, max_examples)}.{fallback}"))
        if blank:
            issues.append(Issue(
                "warning",
                f"Field '{key}': {len(blank)} value(s) of '{primary}' are found but the looked-up "
                f"cell is empty — {{{{{placeholder}}}}} will be empty for "
                f"{sum(len(r) for _, r in blank)} row(s). e.g. {_examples(blank, [], max_examples)}."))
    return issues


def _examples(items, sheet_keys, limit: int) -> str:
    parts = []
    for value, row_numbers in items[:limit]:
        shown = ", ".join(str(n) for n in row_numbers[:3]) + ("…" if len(row_numbers) > 3 else "")
        text = f"'{value}' (row{'s' if len(row_numbers) > 1 else ''} {shown})"
        hint = suggest_name(value, sheet_keys) if sheet_keys else None
        if hint:
            text += f" → did you mean '{hint}'?"
        parts.append(text)
    more = len(items) - limit
    return "; ".join(parts) + (f"; +{more} more" if more > 0 else "")
