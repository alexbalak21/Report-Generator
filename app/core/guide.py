"""
The in-app guide: loading, and the text copied to the clipboard for AI assistants.
No tkinter here.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from app.utils.paths import get_app_resource_path

GUIDE_FILE = "guide.md"

_HIDDEN_CONFIG_KEYS = ("data_file", "template_file", "output_dir")
_COMMENT_LINE = re.compile(r"^\s*<!--.*?-->\s*$")

AI_PREAMBLE = """\
You are helping me use "Report Generator", a desktop app that fills a Word template (.docx) with one
row of an Excel file (.xlsx), following a JSON mapping file. Read the documentation below, then the
description of MY profile, and help me fix or improve my mapping.json, my Word template or my Excel
data. Follow section 0 of the documentation (never invent names, return complete valid JSON, explain
each change in one line, ask when something is missing).
"""

AI_CLOSING = """\

---
My question / problem:
"""


def load_guide() -> str:
    """Guide text without the HTML comment lines used for maintenance."""
    path = Path(get_app_resource_path(GUIDE_FILE))
    text = path.read_text(encoding="utf-8")
    return "\n".join(line for line in text.splitlines() if not _COMMENT_LINE.match(line)).strip() + "\n"


def headings(text: str, level: int = 2) -> list[tuple[str, int]]:
    """[(title, 1-based line number)] of the headings of the given level (outside code blocks)."""
    found, in_code = [], False
    marker = "#" * level + " "
    for number, line in enumerate(text.splitlines(), start=1):
        if line.startswith("```"):
            in_code = not in_code
        elif not in_code and line.startswith(marker):
            found.append((line[len(marker):].strip(), number))
    return found


def _sheet_summary(data_path: str) -> list[str]:
    """Header names and row counts per sheet. Cell values are never included."""
    import openpyxl
    lines = []
    wb = openpyxl.load_workbook(data_path, read_only=True, data_only=True)
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [str(c).strip() for c in first if c is not None]
            rows = max((ws.max_row or 1) - 1, 0)
            lines.append(f'- Sheet "{name}" ({rows} data row(s)): ' +
                         (" | ".join(f'"{h}"' for h in headers) if headers else "(no headers)"))
    finally:
        wb.close()
    return lines


def build_profile_context(profile, issues) -> str:
    """
    Markdown description of the active profile for an AI: check results, mapping.json,
    template placeholders, Excel headers. No cell values; the folder paths stored in the
    mapping's config block are hidden.
    """
    out = ["# MY PROFILE (read from my Report Generator)", "", f"Profile name: {profile.name}", ""]

    out += ["## Check results shown by the app", ""]
    if issues:
        out += [("- ERROR: " if i.severity == "error" else "- Warning: ") + i.message for i in issues]
    else:
        out.append("- No problem found.")

    out += ["", "## mapping.json", ""]
    try:
        raw = Path(profile.mapping_path).read_text(encoding="utf-8")
        try:
            data = json.loads(raw)
            config = data.get("config") if isinstance(data, dict) else None
            if isinstance(config, dict):
                for key in _HIDDEN_CONFIG_KEYS:
                    if key in config:
                        config[key] = "(path hidden)"
            raw = json.dumps(data, indent=2, ensure_ascii=False)
        except ValueError:
            out.append("(The file is not valid JSON; here is its raw content.)")
        out += ["```json", raw.rstrip(), "```"]
    except OSError as exc:
        out.append(f"(Could not read mapping.json: {exc})")

    out += ["", "## Placeholders found in my Word template", ""]
    try:
        from app.core.word_processor import WordProcessor
        names = sorted(set(WordProcessor(profile.template_path).extract_placeholders()))
        out.append(", ".join("{{" + n + "}}" for n in names) if names else "(none found)")
    except Exception as exc:
        out.append(f"(Could not read template.docx: {exc})")

    out += ["", "## Excel data file — headers only (cell values are not included)", ""]
    try:
        out += _sheet_summary(profile.data_path)
    except Exception as exc:
        out.append(f"(Could not read data.xlsx: {exc})")

    return "\n".join(out) + "\n"


def build_ai_message(guide: str, context: str | None = None) -> str:
    """Text copied by 'Copy guide + my profile': instruction + guide + profile + question stub."""
    parts = [AI_PREAMBLE, "=" * 60, guide.rstrip()]
    if context:
        parts += ["", "=" * 60, context.rstrip()]
    parts.append(AI_CLOSING)
    return "\n".join(parts)
