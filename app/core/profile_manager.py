"""
Profiles: a profile bundles everything needed to produce one kind of report.

    <profiles root>/<slug>/
        mapping.json    field rules + settings (the existing mapping format)
        template.docx   Word template with {{placeholders}}
        data.xlsx       Excel data (data sheet + lookup sheets)

No tkinter in here, so everything is unit-testable.
"""
from __future__ import annotations

import datetime
import json
import os
import re
import shutil
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from app.core import profile_checks
from app.core.mapping_loader import MappingLoader
from app.core.profile_checks import Issue, placeholder_name
from app.repository import config_repository as repo
from app.utils.paths import get_profiles_dir, get_data_path

MAPPING_FILE  = "mapping.json"
TEMPLATE_FILE = "template.docx"
DATA_FILE     = "data.xlsx"
PROFILE_FILES = (MAPPING_FILE, TEMPLATE_FILE, DATA_FILE)

KEY_ACTIVE_PROFILE     = "active_profile"
KEY_PROFILES_MIGRATED  = "profiles_migrated"

# Keys that used to hold absolute paths; a profile's files always live in its folder.
_PATH_KEYS = ("data_file", "template_file")

_WINDOWS_RESERVED = {"con", "prn", "aux", "nul",
                     *(f"com{i}" for i in range(1, 10)),
                     *(f"lpt{i}" for i in range(1, 10))}


class ProfileError(Exception):
    """Raised for user-facing profile problems (bad name, missing file, …)."""


@dataclass
class Profile:
    slug: str
    name: str
    folder: Path
    next_line: int = 2
    output_dir: str = ""

    @property
    def mapping_path(self) -> str:
        return str(self.folder / MAPPING_FILE)

    @property
    def template_path(self) -> str:
        return str(self.folder / TEMPLATE_FILE)

    @property
    def data_path(self) -> str:
        return str(self.folder / DATA_FILE)


def slugify(name: str) -> str:
    """Folder-safe slug: 'Fish quality (v2)' -> 'fish-quality-v2'."""
    text = unicodedata.normalize("NFKD", name or "")
    text = text.encode("ascii", "ignore").decode("ascii").lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    if not text:
        text = "profile"
    if text in _WINDOWS_RESERVED:
        text += "-profile"
    return text


class ProfileManager:
    def __init__(self, root: str | os.PathLike | None = None):
        self._root_override = Path(root) if root else None

    # ------------------------------------------------------------------
    # Basics
    # ------------------------------------------------------------------

    @property
    def root(self) -> Path:
        root = self._root_override or Path(get_profiles_dir())
        root.mkdir(parents=True, exist_ok=True)
        return root

    def _to_profile(self, row: dict) -> Profile:
        return Profile(
            slug=row["slug"],
            name=row["name"],
            folder=Path(row["folder"]),
            next_line=max(2, int(row["next_line"] or 2)),
            output_dir=row["output_dir"] or "",
        )

    def list(self) -> list[Profile]:
        self.scan()
        return [self._to_profile(r) for r in repo.profile_list()]

    def get(self, slug: str) -> Profile:
        row = repo.profile_get(slug)
        if row is None:
            raise ProfileError(f"Unknown profile: {slug}")
        return self._to_profile(row)

    def find_by_name(self, name: str) -> Profile | None:
        wanted = name.strip().lower()
        for row in repo.profile_list():
            if row["name"].strip().lower() == wanted:
                return self._to_profile(row)
        return None

    def active_slug(self) -> str | None:
        slug = repo.config_get(KEY_ACTIVE_PROFILE)
        if slug and repo.profile_get(slug):
            return slug
        return None

    def set_active(self, slug: str) -> None:
        repo.config_set(KEY_ACTIVE_PROFILE, slug)
        repo.profile_update(slug, last_used=datetime.datetime.now().isoformat())

    def set_next_line(self, slug: str, line: int) -> None:
        repo.profile_update(slug, next_line=max(2, int(line)))

    def set_output_dir(self, slug: str, directory: str) -> None:
        repo.profile_update(slug, output_dir=directory)

    # ------------------------------------------------------------------
    # Names / slugs
    # ------------------------------------------------------------------

    def _check_name(self, name: str, ignore_slug: str | None = None) -> str:
        name = (name or "").strip()
        if not name:
            raise ProfileError("Please enter a profile name.")
        if len(name) > 80:
            raise ProfileError("Profile name is too long (80 characters max).")
        other = self.find_by_name(name)
        if other and other.slug != ignore_slug:
            raise ProfileError(f"A profile named '{other.name}' already exists.")
        return name

    def _unique_slug(self, name: str) -> str:
        base = slugify(name)
        slug, n = base, 2
        while repo.profile_get(slug) or (self.root / slug).exists():
            slug = f"{base}-{n}"
            n += 1
        return slug

    # ------------------------------------------------------------------
    # Folder scan (register hand-copied folders, forget deleted ones)
    # ------------------------------------------------------------------

    def scan(self) -> None:
        known = {r["slug"]: r for r in repo.profile_list()}

        for slug, row in known.items():
            if not Path(row["folder"]).is_dir():
                repo.profile_delete(slug)

        for entry in sorted(self.root.iterdir()):
            if not entry.is_dir() or entry.name.startswith(("_", ".")):
                continue
            if entry.name in known or not (entry / MAPPING_FILE).is_file():
                continue
            try:
                name = self._read_profile_name(entry / MAPPING_FILE) or entry.name
                name = self._free_name(name)
                repo.profile_add(entry.name, name, str(entry))
            except Exception:
                continue

    def _free_name(self, name: str) -> str:
        candidate, n = name, 2
        while self.find_by_name(candidate):
            candidate = f"{name} ({n})"
            n += 1
        return candidate

    @staticmethod
    def _read_profile_name(mapping_file: Path) -> str:
        try:
            with open(mapping_file, "r", encoding="utf-8") as f:
                raw = json.load(f)
            return str(raw.get("config", {}).get("profile_name", "")).strip()
        except Exception:
            return ""

    # ------------------------------------------------------------------
    # Create / duplicate / rename / delete
    # ------------------------------------------------------------------

    def create(self, name: str, mapping: str | None = None,
               template: str | None = None, data: str | None = None,
               output_dir: str = "") -> Profile:
        """
        Create a profile. Each source file is optional: a missing mapping becomes
        an empty starter mapping; a missing template/data file is flagged by
        validate() until the user adds one.
        """
        name = self._check_name(name)
        slug = self._unique_slug(name)
        folder = self.root / slug
        folder.mkdir(parents=True)
        try:
            self._install_mapping(folder / MAPPING_FILE, mapping, name)
            if template:
                self._copy_file(template, folder / TEMPLATE_FILE)
            if data:
                self._copy_file(data, folder / DATA_FILE)
            repo.profile_add(slug, name, str(folder), output_dir=output_dir or None)
        except Exception:
            shutil.rmtree(folder, ignore_errors=True)
            raise
        return self.get(slug)

    def create_from(self, name: str, source_slug: str | None = None,
                    mapping: str | None = None, template: str | None = None,
                    data: str | None = None) -> Profile:
        """create(), optionally starting from a copy of another profile (files given override it)."""
        if not source_slug:
            return self.create(name, mapping=mapping, template=template, data=data)
        profile = self.duplicate(source_slug, name)
        try:
            if mapping:
                self.replace_mapping(profile.slug, mapping)
            if template:
                self.replace_template(profile.slug, template)
            if data:
                self.replace_data(profile.slug, data)
        except Exception:
            shutil.rmtree(profile.folder, ignore_errors=True)
            repo.profile_delete(profile.slug)
            raise
        return self.get(profile.slug)

    def duplicate(self, slug: str, new_name: str) -> Profile:
        src = self.get(slug)
        new_name = self._check_name(new_name)
        new_slug = self._unique_slug(new_name)
        folder = self.root / new_slug
        folder.mkdir(parents=True)
        try:
            for filename in (TEMPLATE_FILE, DATA_FILE):
                if (src.folder / filename).is_file():
                    shutil.copy2(src.folder / filename, folder / filename)
            self._install_mapping(folder / MAPPING_FILE,
                                  src.mapping_path if Path(src.mapping_path).is_file() else None,
                                  new_name)
            repo.profile_add(new_slug, new_name, str(folder), output_dir=src.output_dir or None)
        except Exception:
            shutil.rmtree(folder, ignore_errors=True)
            raise
        return self.get(new_slug)

    def rename(self, slug: str, new_name: str) -> Profile:
        profile = self.get(slug)
        new_name = self._check_name(new_name, ignore_slug=slug)
        repo.profile_update(slug, name=new_name)
        try:  # keep the hand-copied-folder case working
            loader = MappingLoader(profile.mapping_path)
            loader.update_config(profile_name=new_name)
        except Exception:
            pass
        return self.get(slug)

    def delete(self, slug: str) -> Path | None:
        """Move the profile folder to <root>/_trash/ (never hard-deletes). Returns the new path."""
        profile = self.get(slug)
        trashed = None
        if profile.folder.is_dir():
            trash = self.root / "_trash"
            trash.mkdir(exist_ok=True)
            stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
            trashed = trash / f"{slug}-{stamp}"
            shutil.move(str(profile.folder), str(trashed))
        repo.profile_delete(slug)
        if repo.config_get(KEY_ACTIVE_PROFILE) == slug:
            repo.config_set(KEY_ACTIVE_PROFILE, "")
        return trashed

    # ------------------------------------------------------------------
    # Replace files
    # ------------------------------------------------------------------

    def replace_data(self, slug: str, src: str) -> None:
        self._replace_file(slug, src, DATA_FILE, (".xlsx", ".xlsm"))

    def replace_template(self, slug: str, src: str) -> None:
        self._replace_file(slug, src, TEMPLATE_FILE, (".docx",))

    def replace_mapping(self, slug: str, src: str) -> None:
        profile = self.get(slug)
        dest = profile.folder / MAPPING_FILE
        if dest.is_file():
            shutil.copy2(dest, str(dest) + ".bak")
        self._install_mapping(dest, src, profile.name)

    def _replace_file(self, slug: str, src: str, filename: str, extensions: tuple) -> None:
        if Path(src).suffix.lower() not in extensions:
            raise ProfileError(f"Expected a {' / '.join(extensions)} file.")
        profile = self.get(slug)
        dest = profile.folder / filename
        if Path(src).resolve() == dest.resolve():
            return
        if dest.is_file():
            shutil.copy2(dest, str(dest) + ".bak")
        self._copy_file(src, dest)

    # ------------------------------------------------------------------
    # Mapping install: copy as JSON, drop absolute paths, stamp profile name
    # ------------------------------------------------------------------

    def _install_mapping(self, dest: Path, src: str | None, name: str) -> None:
        if src:
            if not os.path.isfile(src):
                raise ProfileError(f"Mapping file not found: {src}")
            raw = MappingLoader(src).load_raw()
        else:
            raw = {"config": {"date_format": "%d/%m/%Y"}}
        if not isinstance(raw, dict):
            raise ProfileError("The mapping file must contain a JSON object.")
        config = raw.get("config")
        if not isinstance(config, dict):
            config = {}
        for key in _PATH_KEYS:
            config.pop(key, None)
        config["profile_name"] = name
        raw = {"config": config, **{k: v for k, v in raw.items() if k != "config"}}
        with open(dest, "w", encoding="utf-8") as f:
            json.dump(raw, f, indent=2, ensure_ascii=False)

    @staticmethod
    def _copy_file(src: str, dest: Path) -> None:
        if not os.path.isfile(src):
            raise ProfileError(f"File not found: {src}")
        shutil.copy2(src, dest)

    def file_info(self, slug: str) -> dict[str, dict]:
        """{filename: {"exists", "size", "modified"}} for the three profile files."""
        profile = self.get(slug)
        info = {}
        for filename in PROFILE_FILES:
            path = profile.folder / filename
            if path.is_file():
                st = path.stat()
                info[filename] = {
                    "exists": True, "size": st.st_size,
                    "modified": datetime.datetime.fromtimestamp(st.st_mtime),
                }
            else:
                info[filename] = {"exists": False, "size": 0, "modified": None}
        return info

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------

    def validate(self, slug: str, deep: bool = True) -> list[Issue]:
        """
        Check a profile. deep=True also runs every lookup rule over the data rows
        (finds values that would leave a placeholder empty); it reads the whole
        workbook, so pass deep=False for a quick structural check.
        """
        issues: list[Issue] = []
        profile = self.get(slug)

        if not profile.folder.is_dir():
            return [Issue("error", f"Profile folder not found: {profile.folder}")]

        mapping_ok = template_ok = data_ok = False
        mapping: dict = {}

        # mapping
        if not Path(profile.mapping_path).is_file():
            issues.append(Issue("error", f"{MAPPING_FILE} is missing."))
        else:
            try:
                loader = MappingLoader(profile.mapping_path)
                mapping = loader.load()
                config = loader.load_config()
                if not isinstance(mapping, dict) or not isinstance(config, dict):
                    raise ValueError("expected a JSON object")
                mapping_ok = True
            except Exception as exc:
                issues.append(Issue("error", f"{MAPPING_FILE} is not valid: {exc}"))
                config = {}

        # template
        if not Path(profile.template_path).is_file():
            issues.append(Issue("error", f"{TEMPLATE_FILE} is missing."))
        else:
            template_ok = True

        # data
        if not Path(profile.data_path).is_file():
            issues.append(Issue("error", f"{DATA_FILE} is missing."))
        else:
            data_ok = True

        headers: set[str] | None = None
        if data_ok:
            headers, sheet_issue = self._read_headers(
                profile.data_path, config.get("data_sheet") if mapping_ok else None)
            if sheet_issue:
                issues.append(sheet_issue)

        placeholders: set[str] | None = None
        if template_ok:
            try:
                from app.core.word_processor import WordProcessor
                placeholders = set(WordProcessor(profile.template_path).extract_placeholders())
            except Exception as exc:
                issues.append(Issue("error", f"{TEMPLATE_FILE} cannot be read: {exc}"))

        if mapping_ok:
            for key, rule in mapping.items():
                if key == "file_name" or not isinstance(rule, dict) or "column" not in rule:
                    continue
                column = str(rule.get("column", "")).strip()
                if headers is not None and column and column not in headers:
                    hint = profile_checks._did_you_mean(column, headers)
                    issues.append(Issue(
                        "warning", f"Field '{key}': column '{column}' not found in the data sheet.{hint}"))
                name = placeholder_name(rule.get("placeholder"))
                if placeholders is not None and name and name not in placeholders:
                    issues.append(Issue(
                        "warning", f"Field '{key}': placeholder {{{{{name}}}}} "
                                   f"not found in the template."))

            issues.extend(profile_checks.check_operations(mapping))
            issues.extend(profile_checks.check_file_name(mapping, headers))
            if placeholders is not None:
                issues.extend(profile_checks.check_placeholder_coverage(mapping, placeholders))

            if data_ok and headers is not None:
                try:
                    sheets = profile_checks.read_all_headers(profile.data_path)
                    issues.extend(profile_checks.check_lookup_structure(mapping, headers, sheets))
                    if deep:
                        issues.extend(profile_checks.check_lookup_data(
                            mapping, config, profile.data_path))
                except Exception as exc:
                    issues.append(Issue("warning", f"Could not check the lookups: {exc}"))
        return issues

    @staticmethod
    def _read_headers(data_path: str, data_sheet: str | None):
        """Return (header set | None, Issue | None) without loading the whole workbook."""
        import openpyxl
        try:
            wb = openpyxl.load_workbook(data_path, read_only=True, data_only=True)
        except Exception as exc:
            return None, Issue("error", f"{DATA_FILE} cannot be opened: {exc}")
        try:
            names = wb.sheetnames
            sheet_name = None
            if data_sheet:
                sheet_name = next((n for n in names if n == data_sheet), None) or \
                             next((n for n in names if n.lower() == str(data_sheet).lower()), None)
                if sheet_name is None:
                    return None, Issue(
                        "error", f"Data sheet '{data_sheet}' not found in {DATA_FILE} "
                                 f"(sheets: {', '.join(names)}).")
            elif names:
                sheet_name = names[0]
            if sheet_name is None:
                return None, Issue("error", f"{DATA_FILE} has no sheets.")
            first_row = next(wb[sheet_name].iter_rows(min_row=1, max_row=1, values_only=True), ())
            return {str(c).strip() for c in first_row if c is not None}, None
        finally:
            wb.close()

    # ------------------------------------------------------------------
    # Export / import (.zip)
    # ------------------------------------------------------------------

    def export_zip(self, slug: str, dest: str) -> None:
        profile = self.get(slug)
        with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zf:
            for filename in PROFILE_FILES:
                path = profile.folder / filename
                if path.is_file():
                    zf.write(path, filename)

    def import_zip(self, src: str, name: str | None = None) -> Profile:
        if not zipfile.is_zipfile(src):
            raise ProfileError("Not a valid profile archive (.zip).")
        with zipfile.ZipFile(src) as zf:
            members = {Path(n).name: n for n in zf.namelist()
                       if not n.endswith("/") and Path(n).name in PROFILE_FILES}
            if MAPPING_FILE not in members:
                raise ProfileError(f"The archive does not contain {MAPPING_FILE}.")
            tmp = self.root / "_import_tmp"
            shutil.rmtree(tmp, ignore_errors=True)
            tmp.mkdir(parents=True)
            try:
                # Extract by fixed file name only (no path from the archive is ever used).
                for filename, member in members.items():
                    with zf.open(member) as fin, open(tmp / filename, "wb") as fout:
                        shutil.copyfileobj(fin, fout)
                display = (name or self._read_profile_name(tmp / MAPPING_FILE)
                           or Path(src).stem)
                display = self._free_name(display.strip() or "Imported profile")
                return self.create(
                    display,
                    mapping=str(tmp / MAPPING_FILE),
                    template=str(tmp / TEMPLATE_FILE) if TEMPLATE_FILE in members else None,
                    data=str(tmp / DATA_FILE) if DATA_FILE in members else None,
                )
            finally:
                shutil.rmtree(tmp, ignore_errors=True)

    # ------------------------------------------------------------------
    # One-time migration from the pre-profiles setup
    # ------------------------------------------------------------------

    def migrate_legacy(self) -> dict:
        """
        Turn the mapping files registered in the old 'mappings' table into
        profiles (copying their data/template files). Runs once. Originals are
        never modified or deleted.

        Returns {"created": [names], "skipped": [(path, reason)]}.
        """
        result = {"created": [], "skipped": []}
        if repo.config_get(KEY_PROFILES_MIGRATED):
            return result
        repo.config_set(KEY_PROFILES_MIGRATED, "1")

        active_mapping = repo.config_get("mapping_path") or repo.config_get("last_mapping_path") or ""
        g_excel = repo.config_get("last_excel_path") or ""
        g_docx  = repo.config_get("last_docx_path") or ""
        g_out   = repo.config_get("last_output_dir") or ""
        g_line  = repo.config_get("last_line_number")
        active_slug = None

        for _, path in repo.mapping_list():
            if Path(path).suffix.lower() not in (".json", ".xlsx"):
                continue
            if not os.path.isfile(path):
                result["skipped"].append((path, "file not found"))
                continue
            try:
                cfg = MappingLoader(path).load_config()
                is_active = os.path.abspath(path) == os.path.abspath(active_mapping) if active_mapping else False

                data = cfg.get("data_file") or ""
                template = cfg.get("template_file") or ""
                if is_active:
                    data = data if os.path.isfile(data) else g_excel
                    template = template if os.path.isfile(template) else g_docx
                data = data if os.path.isfile(data) else None
                template = template if os.path.isfile(template) else None

                base = str(cfg.get("profile_name") or Path(path).stem).strip() or "Profile"
                profile = self.create(
                    self._free_name(base), mapping=path, template=template, data=data,
                    output_dir=(g_out if is_active and g_out else cfg.get("output_dir", "")) or "",
                )
                if is_active:
                    active_slug = profile.slug
                    if g_line and g_line.isdigit():
                        self.set_next_line(profile.slug, int(g_line) + 1)
                result["created"].append(profile.name)
            except Exception as exc:
                result["skipped"].append((path, str(exc)))

        if active_slug:
            repo.config_set(KEY_ACTIVE_PROFILE, active_slug)

        if not repo.profile_list():
            result["created"].extend(self._seed_defaults())
        return result

    def _seed_defaults(self) -> list[str]:
        """Fresh install: offer the sample reports shipped in the /data folder."""
        samples = [
            ("Rapport",     "rapport_mapping.json",         "rapport_template.docx",     "Rapport_data.xlsx"),
            ("New rapport", "new_rapport_mapping_new.json", "new_rapport_template.docx", "new_data.xlsx"),
        ]
        created = []
        for name, mapping, template, data in samples:
            m, t, d = (get_data_path(x) for x in (mapping, template, data))
            if not os.path.isfile(m):
                continue
            try:
                profile = self.create(name, mapping=m,
                                      template=t if os.path.isfile(t) else None,
                                      data=d if os.path.isfile(d) else None)
                created.append(profile.name)
            except Exception:
                continue
        return created
